"""
tim_easy.py
===========
Genera el TIM "Devolución Logística Inversa" de Easy a partir de números de
entrega. Fuente: BASE DE DATOS WMS (fallback BASEDATOSBEETRACK y BASE DATOS OPL).

Por cada entrega: agrupa por SKU sumando 'Unid'. Una fila por (entrega, SKU).
Solo se rellenan columnas derivables; Motivo/Tipo/Observaciones quedan en blanco.

API:
    build_easy(ids, fecha) -> dict con:
        'xlsx_bytes', 'preview' (list[dict]), 'no_encontrados' (list[str])
"""

from __future__ import annotations

import datetime as dt
import re
from io import BytesIO

import openpyxl

from tools.db_sources import read_beetrack, read_opl, read_wms
from tools.tim_common import TEMPLATES_DIR, copy_row_style, find_header_row

TEMPLATE = TEMPLATES_DIR / "tim_easy_template.xlsx"

# Columnas del formato (1-based) — fila de datos
COL_N = 1        # A  N°
COL_DOC = 2      # B  DOCUMENTO ASOCIADO (entrega)
COL_COD = 6      # F  COD. PRODUCTO (SKU)
COL_DESC = 7     # G  DESCRIPCION
COL_UN = 8       # H  UN
COL_RANGE = (1, 11)  # A..K para copiar estilos

# Anchos reales de la plantilla (se re-aseguran por si Excel ignora los heredados)
ANCHOS = {"A": 4.4, "B": 37.4, "C": 21.3, "D": 14.0, "E": 0.1, "F": 24.7,
          "G": 64.7, "H": 6.3, "I": 15.9, "J": 73.1, "K": 26.3}


def _aggregate(ids: list[str]):
    """Devuelve (filas, no_encontrados). filas: [{'entrega','sku','desc','un'}]."""
    wms = read_wms()
    beet = read_beetrack()
    # Tercera fuente: los pedidos de OPL. El 'Id entrega' (col C de BASE DATOS OPL) no
    # existe en WMS ni en Beetrack, asi que pegarlo daba "sin datos en las bases"
    # aunque la SubOrden del MISMO pedido si se encontrara (22-10).
    opl = read_opl()
    filas: list[dict] = []
    no_encontrados: list[str] = []

    for ent in ids:
        # OPL antes que Beetrack: para un pedido de OPL, Beetrack trae el dato
        # derivado (SKU vacio o con el otro codigo, y cantidad "0 / 1" que da 0),
        # mientras que OPL tiene el SKU, la descripcion y las unidades de origen.
        lineas = wms.get(ent) or opl.get(ent) or beet.get(ent)
        if not lineas:
            no_encontrados.append(ent)
            continue
        # agrupar por SKU sumando unidades; conservar primera descripción vista
        agg: dict[str, dict] = {}
        orden_sku: list[str] = []
        for ln in lineas:
            sku = ln["sku"]
            if sku not in agg:
                agg[sku] = {"desc": ln["desc"], "un": 0}
                orden_sku.append(sku)
            agg[sku]["un"] += ln.get("unid", 0) or 0
            if not agg[sku]["desc"] and ln["desc"]:
                agg[sku]["desc"] = ln["desc"]
        for sku in orden_sku:
            un = agg[sku]["un"]
            filas.append({
                "entrega": ent,
                "sku": sku,
                "desc": agg[sku]["desc"],
                "un": int(un) if float(un).is_integer() else un,
            })
    return filas, no_encontrados


# Un retiro (cambio anulado) trae en su propia descripción el N° de la entrega
# original entre paréntesis, ej. "CAMBIO COMEDOR 8S CLARICE 136X136X77 (2918619547)".
_REF_RE = re.compile(r"\((\d{6,})\)\s*$")


def _ref_entrega(desc: str) -> str | None:
    """N° de entrega referenciado al final de la descripción de un retiro, si hay."""
    m = _REF_RE.search(desc or "")
    return m.group(1) if m else None


# Palabra suelta: "RECAMBIO FILTRO AGUA" es un producto real, no un retiro.
# Verificado contra las 17.949 descripciones del Sheet: clasifica igual que el
# substring anterior (0 diferencias), pero ya no puede confundirse.
_RETIRO_RE = re.compile(r"\b(CAMBIO|RETIRO)S?\b")


def _es_retiro(desc: str) -> bool:
    # ponytail: heurístico por texto en la descripción (así es como el equipo ya
    # identifica un cambio/retiro a simple vista). Techo: no exige el paréntesis
    # con N° al final, así un "RETIRO CD JAULA" sin N° igual cuenta como retiro
    # (queda suelto en su propia fila, que es lo correcto).
    return bool(_RETIRO_RE.search((desc or "").upper()))


def _merge_cambios(filas: list[dict]) -> list[dict]:
    """Fusiona cada retiro (cambio anulado) con la fila de su entrega original.

    Se emparejan solo cuando la descripción del retiro referencia OTRA entrega que
    también está en `filas` (evita falsos positivos: un retiro que se referencia a
    sí mismo, o un "RETIRO CD ..." sin N° de entrega, quedan sueltos como antes).
    La entrega conserva TODAS sus filas (una por SKU); a cada una se le agrega el N° del
    retiro en 'entrega'. Solo cuando la entrega trae un único SKU se le copia la
    descripción del retiro (que es del mismo producto y aporta el "CAMBIO"/"RETIRO" + el
    N° entre paréntesis); con varios SKU se conserva el nombre de cada producto, si no se
    perdería uno y se duplicaría otro.

    La clasificación retiro/entrega se hace UNA sola vez antes de mutar nada: al fusionar
    se sobreescribe la desc del destino con la del retiro, así que reclasificar sobre la
    marcha haría que el destino pareciera retiro, se fusionara consigo mismo y su fila se
    perdiera (pasaba cuando el retiro venía antes que su entrega en la lista).
    """
    clasificadas = [(f, _es_retiro(f["desc"])) for f in filas]
    por_entrega: dict[str, list[dict]] = {}
    for f, retiro in clasificadas:
        if not retiro:
            por_entrega.setdefault(f["entrega"], []).append(f)
    resultado: list[dict] = []
    for f, retiro in clasificadas:
        if not retiro:
            resultado.append(f)
            continue
        ref = _ref_entrega(f["desc"])
        objetivo = por_entrega.get(ref) if ref and ref != f["entrega"] else None
        if not objetivo:
            resultado.append(f)  # retiro huérfano/auto-referenciado: queda solo
            continue
        for fila in objetivo:
            fila["entrega"] = f"{fila['entrega']} - {f['entrega']}"
        if len(objetivo) == 1:
            objetivo[0]["desc"] = f["desc"]
    return resultado


def build_easy(ids: list[str], fecha: dt.date | None = None) -> dict:
    fecha = fecha or dt.date.today()
    filas, no_encontrados = _aggregate(ids)
    n_antes = len(filas)
    filas = _merge_cambios(filas)

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active

    # Normalizar la vista del libro: la plantilla heredó firstSheet=1 (de un original
    # con 2 hojas) pero ahora solo hay 1 hoja → índice fuera de rango → Excel pide
    # reparar y al recuperar descarta anchos/formato. Forzar a la única hoja (índice 0).
    for view in getattr(wb, "views", []) or []:
        view.firstSheet = 0
        view.activeTab = 0

    # Re-asegurar anchos de columna (Excel a veces ignora los heredados).
    # Asignar .width ya marca customWidth=True automáticamente en openpyxl.
    for col, w in ANCHOS.items():
        ws.column_dimensions[col].width = w

    # Fecha (B6)
    ws["B6"] = fecha.strftime("%d-%m-%Y")

    header_row = find_header_row(ws, "DOCUMENTO ASOCIADO")  # fila 8
    data_start = header_row + 1

    # Leyenda "Motivo registro": buscar SOLO en la columna A y debajo del encabezado
    # (el header de tabla C8='MOTIVO REGISTRO' también coincidiría con un match global).
    legend_row = _find_legend_row(ws, "Motivo registro", start=data_start)
    last_template_row = legend_row - 2  # deja una fila en blanco antes de la leyenda

    # Si hay más filas que el espacio pre-formateado, insertar filas ANTES de la
    # leyenda para que ésta (y el bloque de firmas) se desplacen sin sobrescribirse.
    capacidad = last_template_row - data_start + 1
    if len(filas) > capacidad:
        faltan = len(filas) - capacidad
        ws.insert_rows(last_template_row + 1, faltan)
        for k in range(faltan):
            copy_row_style(ws, data_start, last_template_row + 1 + k,
                           COL_RANGE[0], COL_RANGE[1])
        last_template_row += faltan

    for i, f in enumerate(filas):
        r = data_start + i
        if r > last_template_row:
            copy_row_style(ws, data_start, r, COL_RANGE[0], COL_RANGE[1])
        ws.cell(r, COL_N).value = i + 1
        cB = ws.cell(r, COL_DOC)
        cB.value = _num_or_str(f["entrega"])
        cB.number_format = "0"
        cF = ws.cell(r, COL_COD)
        cF.value = _num_or_str(f["sku"])
        cF.number_format = "0"
        ws.cell(r, COL_DESC).value = f["desc"]
        cH = ws.cell(r, COL_UN)
        cH.value = f["un"]
        cH.number_format = "0"

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return {
        "xlsx_bytes": buf.getvalue(),
        "preview": [
            {"N°": i + 1, "DOCUMENTO ASOCIADO": f["entrega"], "COD. PRODUCTO": f["sku"],
             "DESCRIPCION": f["desc"], "UN": f["un"]}
            for i, f in enumerate(filas)
        ],
        "no_encontrados": no_encontrados,
        "fusionados": n_antes - len(filas),  # retiros absorbidos por su entrega original
    }


def _find_legend_row(ws, marker: str, start: int = 1) -> int:
    """Fila (1-based) donde la columna A == marker, buscando desde `start` hacia abajo."""
    target = marker.strip().lower()
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None and str(v).strip().lower() == target:
            return r
    raise ValueError(f"No se encontró la leyenda '{marker}' en la columna A.")


def _num_or_str(s: str):
    """Convierte a int si es numérico puro (para que Excel lo trate como número)."""
    s = str(s).strip()
    if s.isdigit():
        try:
            return int(s)
        except ValueError:
            return s
    return s
